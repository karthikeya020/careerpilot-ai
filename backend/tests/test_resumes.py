import io

from docx import Document


def _register_and_auth(client, email="resume@example.com"):
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Password1", "full_name": "Resume Student"},
    )
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _build_docx() -> bytes:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Aspiring backend engineer with a passion for distributed systems.")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, SQL, Git, Docker")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("Built a REST API using FastAPI and PostgreSQL for a university course project.")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.Tech Computer Science, 2026")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_upload_docx_resume_parses_sections_and_skills(client) -> None:
    headers = _register_and_auth(client)
    docx_bytes = _build_docx()

    response = client.post(
        "/api/v1/resumes",
        headers=headers,
        files={"file": ("resume.docx", docx_bytes, DOCX_CONTENT_TYPE)},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["parsing_status"] == "parsed"

    section_types = {s["section_type"] for s in body["sections"]}
    assert {"summary", "skills", "projects", "education"}.issubset(section_types)

    skill_names = {rs["skill"]["name"] for rs in body["resume_skills"]}
    assert {"Python", "SQL", "Git", "Docker", "FastAPI", "PostgreSQL"}.issubset(skill_names)

    projects_evidence = next(rs for rs in body["resume_skills"] if rs["skill"]["name"] == "FastAPI")
    assert "FastAPI" in projects_evidence["evidence_snippet"]

    get_response = client.get("/api/v1/resumes/me", headers=headers)
    assert get_response.status_code == 200
    assert get_response.json()["id"] == body["id"]


def test_upload_rejects_unsupported_file_type(client) -> None:
    headers = _register_and_auth(client)
    response = client.post(
        "/api/v1/resumes",
        headers=headers,
        files={"file": ("resume.exe", b"not a resume", "application/octet-stream")},
    )
    assert response.status_code == 422


def test_upload_rejects_empty_file(client) -> None:
    headers = _register_and_auth(client)
    response = client.post(
        "/api/v1/resumes",
        headers=headers,
        files={"file": ("resume.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 422


def test_analysis_endpoint_before_upload_reports_no_resume(client) -> None:
    headers = _register_and_auth(client, "analysis-empty@example.com")
    response = client.get("/api/v1/resumes/me/analysis", headers=headers)
    assert response.status_code == 200
    assert response.json()["has_resume"] is False


def test_analysis_endpoint_after_upload_returns_bullets_consistency_and_parseability(client) -> None:
    headers = _register_and_auth(client, "analysis-full@example.com")
    document = Document()
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Aspiring backend engineer with a passion for distributed systems.")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, Kubernetes")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("- Built a REST API using Python, cutting response latency by 40%.")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.Tech Computer Science, 2026")
    buffer = io.BytesIO()
    document.save(buffer)

    upload = client.post(
        "/api/v1/resumes", headers=headers, files={"file": ("resume.docx", buffer.getvalue(), DOCX_CONTENT_TYPE)}
    )
    assert upload.status_code == 201

    response = client.get("/api/v1/resumes/me/analysis", headers=headers)
    assert response.status_code == 200
    body = response.json()
    assert body["has_resume"] is True
    assert len(body["bullet_grades"]) >= 1
    assert body["bullet_grades"][0]["strength"] == "strong"
    flagged = {f["skill_name"] for f in body["self_consistency_flags"]}
    assert "Kubernetes" in flagged
    assert body["parseability"]["score"] > 0.5


def test_recruiter_card_self_view_reflects_own_evidence(client) -> None:
    headers = _register_and_auth(client, "recruiter-card-self@example.com")
    empty_card = client.get("/api/v1/resumes/me/recruiter-card", headers=headers)
    assert empty_card.status_code == 200
    assert empty_card.json()["has_resume"] is False

    docx_bytes = _build_docx()
    upload = client.post(
        "/api/v1/resumes", headers=headers, files={"file": ("resume.docx", docx_bytes, DOCX_CONTENT_TYPE)}
    )
    assert upload.status_code == 201

    card = client.get("/api/v1/resumes/me/recruiter-card", headers=headers)
    assert card.status_code == 200
    body = card.json()
    assert body["has_resume"] is True
    assert 0.0 <= body["trust_score"] <= 1.0
    assert "not a hiring recommendation" in body["disclaimer"].lower()


def test_rewrite_suggestions_never_fabricate_evidence_for_missing_skills(client) -> None:
    headers = _register_and_auth(client, "rewrite-suggestions@example.com")
    docx_bytes = _build_docx()
    upload = client.post(
        "/api/v1/resumes", headers=headers, files={"file": ("resume.docx", docx_bytes, DOCX_CONTENT_TYPE)}
    )
    assert upload.status_code == 201

    search = client.get("/api/v1/job-catalog/search?company=Google", headers=headers)
    assert search.status_code == 200
    listing_id = search.json()[0]["listing"]["id"]

    response = client.get(
        f"/api/v1/resumes/me/rewrite-suggestions?listing_id={listing_id}", headers=headers
    )
    assert response.status_code == 200
    suggestions = response.json()
    assert isinstance(suggestions, list)
    for suggestion in suggestions:
        if not suggestion["has_sufficient_evidence"]:
            assert suggestion["rewritten_bullet"] is None
            assert "no existing evidence" in suggestion["note"].lower()


def test_get_resume_without_upload_returns_404(client) -> None:
    headers = _register_and_auth(client)
    response = client.get("/api/v1/resumes/me", headers=headers)
    assert response.status_code == 404


def _build_docx_with_skills(skills_line: str, projects_line: str) -> bytes:
    document = Document()
    document.add_paragraph("Test Student")
    document.add_paragraph("SKILLS")
    document.add_paragraph(skills_line)
    document.add_paragraph("PROJECTS")
    document.add_paragraph(projects_line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_replacing_a_resume_supersedes_the_old_one_and_removes_its_evidence_from_the_active_profile(client) -> None:
    """DEFECT-001 regression test: uploading Resume B must make Resume A's
    skills disappear from the active Career Twin / JD-match profile, while
    Resume A itself remains in history (not deleted) -- see
    docs/implementation/FINAL_TRUTH_FIRST_DEFECT_LEDGER.md DEFECT-001."""
    headers = _register_and_auth(client, email="staleness@example.com")

    # Resume A: Java-heavy backend resume.
    resume_a_bytes = _build_docx_with_skills(
        "Java, Spring, SQL, Git", "Built a Java Spring Boot microservice backed by a SQL database."
    )
    resp_a = client.post(
        "/api/v1/resumes", headers=headers, files={"file": ("resume-a.docx", resume_a_bytes, DOCX_CONTENT_TYPE)}
    )
    assert resp_a.status_code == 201
    resume_a = resp_a.json()
    assert resume_a["is_active"] is True
    skill_names_a = {rs["skill"]["name"] for rs in resume_a["resume_skills"]}
    assert "Java" in skill_names_a

    # Post a job description requiring Java -- must match while Resume A is active.
    jd_resp = client.post(
        "/api/v1/job-descriptions",
        headers=headers,
        json={
            "title": "Backend Engineer",
            "company": "Acme",
            "raw_text": "We are hiring a backend engineer.\nRequired skills:\nJava\nPandas\n",
        },
    )
    assert jd_resp.status_code == 201
    jd_id = jd_resp.json()["id"]

    match_before = client.get(f"/api/v1/job-descriptions/{jd_id}/match", headers=headers).json()
    matched_before = {s["name"] for s in match_before["matched_skills"]} | {
        s["name"] for s in match_before["partial_skills"]
    }
    assert "Java" in matched_before

    # Resume B: an entirely different Python/data resume, no Java.
    resume_b_bytes = _build_docx_with_skills(
        "Python, Pandas, NumPy, Git", "Built a data pipeline using Python, Pandas, and NumPy."
    )
    resp_b = client.post(
        "/api/v1/resumes", headers=headers, files={"file": ("resume-b.docx", resume_b_bytes, DOCX_CONTENT_TYPE)}
    )
    assert resp_b.status_code == 201
    resume_b = resp_b.json()
    assert resume_b["is_active"] is True
    skill_names_b = {rs["skill"]["name"] for rs in resume_b["resume_skills"]}
    assert "Pandas" in skill_names_b
    assert "Java" not in skill_names_b

    # GET /resumes/me now returns Resume B (the active one).
    active_resp = client.get("/api/v1/resumes/me", headers=headers)
    assert active_resp.json()["id"] == resume_b["id"]

    # Resume A is preserved in history, marked superseded/inactive -- not deleted.
    history_resp = client.get("/api/v1/resumes", headers=headers)
    assert history_resp.status_code == 200
    history = {r["id"]: r for r in history_resp.json()}
    assert resume_a["id"] in history
    assert history[resume_a["id"]]["is_active"] is False
    assert history[resume_a["id"]]["superseded_at"] is not None
    assert history[resume_b["id"]]["is_active"] is True

    # The active JD match no longer counts Java (Resume A's evidence) and now
    # counts Pandas (Resume B's evidence) -- proves recomputation, not staleness.
    match_after = client.get(f"/api/v1/job-descriptions/{jd_id}/match", headers=headers).json()
    matched_after = {s["name"] for s in match_after["matched_skills"]} | {
        s["name"] for s in match_after["partial_skills"]
    }
    missing_after = {s["name"] for s in match_after["missing_skills"]}
    assert "Pandas" in matched_after
    assert "Java" in missing_after
    assert "Java" not in matched_after

    # Reactivating Resume A brings Java back into the active profile.
    reactivate_resp = client.post(f"/api/v1/resumes/{resume_a['id']}/activate", headers=headers)
    assert reactivate_resp.status_code == 200
    assert reactivate_resp.json()["is_active"] is True

    history_after_reactivate = {r["id"]: r for r in client.get("/api/v1/resumes", headers=headers).json()}
    assert history_after_reactivate[resume_a["id"]]["is_active"] is True
    assert history_after_reactivate[resume_b["id"]]["is_active"] is False

    match_reactivated = client.get(f"/api/v1/job-descriptions/{jd_id}/match", headers=headers).json()
    matched_reactivated = {s["name"] for s in match_reactivated["matched_skills"]} | {
        s["name"] for s in match_reactivated["partial_skills"]
    }
    assert "Java" in matched_reactivated


def test_resume_history_and_activation_are_isolated_per_student(client) -> None:
    headers_a = _register_and_auth(client, email="isolation-owner@example.com")
    headers_b = _register_and_auth(client, email="isolation-other@example.com")

    resume_bytes = _build_docx_with_skills("Java, SQL", "A Java project.")
    upload_resp = client.post(
        "/api/v1/resumes", headers=headers_a, files={"file": ("resume.docx", resume_bytes, DOCX_CONTENT_TYPE)}
    )
    resume_id = upload_resp.json()["id"]

    # A different student cannot see it in their own history...
    other_history = client.get("/api/v1/resumes", headers=headers_b).json()
    assert all(r["id"] != resume_id for r in other_history)

    # ...and cannot activate it.
    activate_resp = client.post(f"/api/v1/resumes/{resume_id}/activate", headers=headers_b)
    assert activate_resp.status_code == 404
