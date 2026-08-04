import io

from docx import Document

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

JD_TEXT = """We are looking for a Backend Engineer.

Requirements:
- Bachelor's degree in Computer Science or related field
- 2+ years of experience building backend services
- Strong experience with Python and SQL
- Experience with FastAPI and PostgreSQL
- Familiarity with Docker and Kubernetes
- Demonstrated leadership skills
- Excellent communication skills

Responsibilities:
- Design and build REST APIs
- Collaborate with cross-functional teams
- Mentor junior engineers

Nice to have: experience with GraphQL
"""


def _register_and_auth(client, email="jd@example.com"):
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Password1", "full_name": "JD Student"},
    )
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _upload_resume(client, headers) -> None:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Aspiring backend engineer with a passion for distributed systems and leadership.")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, SQL, Git, Docker")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("Built a REST API using FastAPI and PostgreSQL for a university course project.")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.Tech Computer Science, 2026")
    buffer = io.BytesIO()
    document.save(buffer)

    response = client.post(
        "/api/v1/resumes",
        headers=headers,
        files={"file": ("resume.docx", buffer.getvalue(), DOCX_CONTENT_TYPE)},
    )
    assert response.status_code == 201


def test_add_job_description_extracts_requirements(client) -> None:
    headers = _register_and_auth(client)
    response = client.post(
        "/api/v1/job-descriptions",
        headers=headers,
        json={"title": "Backend Engineer", "company": "Acme", "raw_text": JD_TEXT},
    )
    assert response.status_code == 201
    body = response.json()

    req_types = {r["requirement_type"] for r in body["requirements"]}
    assert {"skill", "experience", "education", "responsibility"}.issubset(req_types)

    skill_reqs = [r for r in body["requirements"] if r["requirement_type"] == "skill"]
    skill_names = {r["skill"]["name"] for r in skill_reqs}
    assert {"Python", "SQL", "FastAPI", "PostgreSQL", "Docker", "Kubernetes"}.issubset(skill_names)

    graphql_req = next(r for r in skill_reqs if r["skill"]["name"] == "GraphQL")
    assert graphql_req["is_required"] is False


def test_match_without_resume_reports_all_missing(client) -> None:
    headers = _register_and_auth(client)
    create_response = client.post(
        "/api/v1/job-descriptions",
        headers=headers,
        json={"title": "Backend Engineer", "company": "Acme", "raw_text": JD_TEXT},
    )
    jd_id = create_response.json()["id"]

    match_response = client.get(f"/api/v1/job-descriptions/{jd_id}/match", headers=headers)
    assert match_response.status_code == 200
    body = match_response.json()
    assert body["matched_skills"] == []
    assert len(body["missing_skills"]) > 0


def test_match_after_resume_upload_classifies_matched_partial_missing(client) -> None:
    headers = _register_and_auth(client)
    _upload_resume(client, headers)

    create_response = client.post(
        "/api/v1/job-descriptions",
        headers=headers,
        json={"title": "Backend Engineer", "company": "Acme", "raw_text": JD_TEXT},
    )
    assert create_response.status_code == 201
    jd_id = create_response.json()["id"]

    match_response = client.get(f"/api/v1/job-descriptions/{jd_id}/match", headers=headers)
    body = match_response.json()

    matched_names = {s["name"] for s in body["matched_skills"]}
    partial_names = {s["name"] for s in body["partial_skills"]}
    missing_names = {s["name"] for s in body["missing_skills"]}

    assert {"Python", "SQL", "FastAPI", "PostgreSQL", "Docker"}.issubset(matched_names)
    assert "Leadership" in partial_names
    assert "Kubernetes" in missing_names
    assert "Communication" in missing_names
    assert 0 < body["coverage"] < 1
    assert body["confidence"] is not None

    # Calling match again must not duplicate evidence or double-count.
    second_response = client.get(f"/api/v1/job-descriptions/{jd_id}/match", headers=headers)
    assert second_response.json() == body


def test_list_and_get_job_description(client) -> None:
    headers = _register_and_auth(client)
    client.post(
        "/api/v1/job-descriptions",
        headers=headers,
        json={"title": "Backend Engineer", "raw_text": JD_TEXT},
    )
    list_response = client.get("/api/v1/job-descriptions", headers=headers)
    assert list_response.status_code == 200
    assert len(list_response.json()) == 1

    jd_id = list_response.json()[0]["id"]
    get_response = client.get(f"/api/v1/job-descriptions/{jd_id}", headers=headers)
    assert get_response.status_code == 200
    assert get_response.json()["title"] == "Backend Engineer"


def test_get_unknown_job_description_404(client) -> None:
    headers = _register_and_auth(client)
    response = client.get(
        "/api/v1/job-descriptions/00000000-0000-0000-0000-000000000000",
        headers=headers,
    )
    assert response.status_code == 404
