import io

from docx import Document

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

JD_TEXT = """We are looking for a Backend Engineer.

Requirements:
- Strong experience with Python and SQL
- Experience with FastAPI and PostgreSQL
- Familiarity with Kubernetes
"""


def _register_and_auth(client, email="dash@example.com"):
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Password1", "full_name": "Dash Student"},
    )
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _onboard(client, headers) -> None:
    response = client.post(
        "/api/v1/onboarding",
        headers=headers,
        json={
            "full_name": "Dash Student",
            "career_goal_description": "Become a backend engineer.",
            "target_role_title": "Backend Engineer",
            "self_assessed_skills": [{"skill_name": "Python", "rating": 0.5}],
        },
    )
    assert response.status_code == 200


def _upload_resume(client, headers) -> None:
    document = Document()
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, SQL")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("Built a service using FastAPI and PostgreSQL.")
    buffer = io.BytesIO()
    document.save(buffer)
    response = client.post(
        "/api/v1/resumes",
        headers=headers,
        files={"file": ("resume.docx", buffer.getvalue(), DOCX_CONTENT_TYPE)},
    )
    assert response.status_code == 201


def test_dashboard_before_any_activity_has_safe_defaults(client) -> None:
    headers = _register_and_auth(client)
    response = client.get("/api/v1/dashboard", headers=headers)
    assert response.status_code == 200
    body = response.json()
    assert body["onboarding_completed"] is False
    assert body["career_twin"] is None
    assert body["mission"] is None
    assert body["resume_status"]["uploaded"] is False
    assert body["job_description_status"]["added"] is False
    assert body["recent_evidence"] == []


def test_dashboard_reflects_full_happy_path(client) -> None:
    headers = _register_and_auth(client)
    _onboard(client, headers)
    _upload_resume(client, headers)

    jd_response = client.post(
        "/api/v1/job-descriptions",
        headers=headers,
        json={"title": "Backend Engineer", "raw_text": JD_TEXT},
    )
    assert jd_response.status_code == 201

    response = client.get("/api/v1/dashboard", headers=headers)
    assert response.status_code == 200
    body = response.json()

    assert body["onboarding_completed"] is True
    assert body["target_role"]["title"] == "Backend Engineer"
    assert body["career_twin"] is not None
    assert body["career_twin"]["version"] >= 3  # onboarding + resume + jd match
    assert body["mission"] is not None
    assert body["resume_status"]["uploaded"] is True
    assert body["resume_status"]["parsing_status"] == "parsed"
    assert body["job_description_status"]["added"] is True
    assert body["job_description_status"]["title"] == "Backend Engineer"
    assert len(body["recent_evidence"]) > 0
    assert len(body["recent_twin_updates"]) >= 3
    assert len(body["recent_audit_events"]) > 0
    assert body["system_trust"]["formula_version"] == "twin-v1"

    twin_response = client.get("/api/v1/career-twin", headers=headers)
    assert twin_response.status_code == 200
    assert twin_response.json()["version"] == body["career_twin"]["version"]

    history_response = client.get("/api/v1/career-twin/history", headers=headers)
    assert len(history_response.json()) == body["career_twin"]["version"]

    active_mission_response = client.get("/api/v1/missions/active", headers=headers)
    assert active_mission_response.status_code == 200
    mission_id = active_mission_response.json()["id"]

    complete_response = client.post(f"/api/v1/missions/{mission_id}/complete", headers=headers)
    assert complete_response.status_code == 200
    assert complete_response.json()["status"] == "completed"

    audit_response = client.get("/api/v1/audit", headers=headers)
    assert audit_response.status_code == 200
    assert len(audit_response.json()) > 0

    skills_response = client.get("/api/v1/skills", headers=headers)
    assert skills_response.status_code == 200
    assert len(skills_response.json()) >= 50
